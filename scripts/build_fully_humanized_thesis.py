import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_section_pagination(section, fmt="decimal", start=None):
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag.endswith('pgNumType'):
            sectPr.remove(child)
    w_ns = nsdecls('w')
    if start is not None:
        xml_str = f'<w:pgNumType {w_ns} w:fmt="{fmt}" w:start="{start}"/>'
    else:
        xml_str = f'<w:pgNumType {w_ns} w:fmt="{fmt}"/>'
    pgNumType = parse_xml(xml_str)
    sectPr.append(pgNumType)

def add_field_code(paragraph, field_text):
    w_ns = nsdecls('w')
    xml_str = (
        f'<w:r {w_ns}><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r {w_ns}><w:instrText xml:space="preserve"> {field_text} </w:instrText></w:r>'
        f'<w:r {w_ns}><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r {w_ns}><w:fldChar w:fldCharType="end"/></w:r>'
    )
    r_elems = parse_xml(f'<w:wrapper {w_ns}>{xml_str}</w:wrapper>')
    for child in list(r_elems):
        paragraph._p.append(child)

def add_footer_page_number(paragraph):
    add_field_code(paragraph, "PAGE")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_content(cell, text, bold=False, font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15, space_after=2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    return run

def add_figure_caption(doc, caption_text):
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    
    r1 = p.add_run("Figure ")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(10)
    r1.font.italic = True

    add_field_code(p, "SEQ Figure \\* ARABIC")

    r2 = p.add_run(f": {caption_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10)
    r2.font.italic = True

def add_table_caption(doc, caption_text):
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    
    r1 = p.add_run("Table ")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(10)
    r1.font.italic = True

    add_field_code(p, "SEQ Table \\* ARABIC")

    r2 = p.add_run(f": {caption_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10)
    r2.font.italic = True

def add_picture_with_caption(doc, img_path, caption_text, width_inches=5.4):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        add_figure_caption(doc, caption_text)
    else:
        print(f"Warning: Image file not found at {img_path}")

print("Setup completed.")
