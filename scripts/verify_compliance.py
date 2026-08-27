import docx
import sys

def verify_doc():
    doc_path = r"d:\FinalProject 2026\vehicle_damade_dash_cloned\BSC-WD-22-36-01.docx"
    doc = docx.Document(doc_path)

    full_text = []
    for p in doc.paragraphs:
        full_text.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                full_text.append(cell.text)

    all_text = " ".join(full_text)
    all_text_lower = all_text.lower()

    forbidden_terms = [
        "gemini",
        "google vision",
        "generative ai vision",
        "api key to damage",
        "gemini api",
        "prone",
    ]

    violations = []
    for term in forbidden_terms:
        if term in all_text_lower:
            violations.append(term)

    print("=" * 60)
    print("THESIS COMPLIANCE AUDIT REPORT")
    print("=" * 60)
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print(f"Total Words Scanned: {len(all_text.split())}")
    
    if violations:
        print(f"FAILED: Found forbidden terms: {violations}")
        for term in violations:
            # find snippets
            idx = 0
            while True:
                idx = all_text_lower.find(term, idx)
                if idx == -1:
                    break
                snippet = all_text[max(0, idx - 50): min(len(all_text), idx + 50)]
                print(f"  - Snippet for '{term}': ...{snippet}...")
                idx += len(term)
    else:
        print("PASSED: Zero forbidden terms found! 100% compliant with negative constraints.")

    required_terms = [
        "W.M.M.G. SENAVIRATHNE",
        "BSC/WD/22/36/01",
        "Bhagya Thilakarathne",
        "Sri Lanka International Buddhist Academy",
        "August 2026",
        "YOLOv8",
        "41.20%",
        "13,945",
    ]

    missing = []
    for r in required_terms:
        if r.lower() not in all_text_lower:
            missing.append(r)

    if missing:
        print(f"FAILED: Missing required metadata: {missing}")
    else:
        print("PASSED: All required academic metadata, supervisor, student, and ML metrics confirmed present!")
    print("=" * 60)

if __name__ == "__main__":
    verify_doc()
